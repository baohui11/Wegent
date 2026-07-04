# SPDX-License-Identifier: Apache-2.0
"""Extract plain text from an uploaded tender file (docx/pdf/txt/md).

Deliberately lightweight (no OCR): docx bodies and text-layer PDFs only.
Scanned PDFs yield no text and are rejected with a clear error; the
knowledge_doc_converter (MinerU) pipeline is the phase-2 upgrade path.
"""

import io
from pathlib import Path

_TEXT_SUFFIXES = {".txt", ".md", ".markdown"}
_MIN_CHARS = 30  # below this we assume a scanned or empty document


class TenderExtractError(ValueError):
    pass


def _decode(content: bytes) -> str:
    for enc in ("utf-8", "gb18030"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _docx_text(content: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def _pdf_text(content: bytes) -> str:
    from PyPDF2 import PdfReader

    reader = PdfReader(io.BytesIO(content))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix in _TEXT_SUFFIXES:
        text = _decode(content)
    elif suffix == ".docx":
        try:
            text = _docx_text(content)
        except Exception as e:
            raise TenderExtractError(f"cannot read docx: {e}") from e
    elif suffix == ".pdf":
        try:
            text = _pdf_text(content)
        except Exception as e:
            raise TenderExtractError(f"cannot read pdf: {e}") from e
    else:
        raise TenderExtractError(
            f"unsupported file type: {suffix or '(none)'}; 支持 .txt/.md/.docx/.pdf"
        )
    if len(text.strip()) < _MIN_CHARS:
        raise TenderExtractError(
            "未能提取到有效文本（可能是扫描版或空文档），请提供文字版"
        )
    return text
