# SPDX-License-Identifier: Apache-2.0
import pytest
from docx import Document

from app.services.bid import assemble_service as assemble
from app.services.bid.parse_pipeline import BidPipelineError
from app.services.bid.workspace import BidWorkspace


def _seed(ws: BidWorkspace) -> None:
    ws.write_json(
        "workspace/tender.json",
        {
            "project": {
                "name": "某信息化项目",
                "id": "ZB-1",
                "budget": "1000000",
                "bid_deadline": "2026-12-31T17:00:00",
            },
            "submission_rules": {"blind_bid": False},
        },
    )
    ws.write_json(
        "workspace/outline.json",
        {"sections": [{"id": "s1", "title": "总述"}, {"id": "s2", "title": "质量"}]},
    )
    ws.path("workspace/sections").mkdir(parents=True, exist_ok=True)
    # Section bodies are title-less: the heading now lives only in the outline,
    # and the assembler injects it (single-document foundation §6 / spike-notes
    # Conclusion B). Previously the body started with `# 总述`.
    ws.path("workspace/sections/s1.md").write_text(
        "本项目由{{bidder}}实施，已通过 CMMI3 评估{{qual:CMMI3}}。",
        encoding="utf-8",
    )
    ws.path("workspace/sections/s2.md").write_text(
        "依托 ISO9001 认证{{qual:ISO9001}}。", encoding="utf-8"
    )
    ws.write_json(
        "corpus/qualifications.json",
        {
            "company": "测试科技有限公司",
            "items": [
                {
                    "id": "CMMI3",
                    "name": "CMMI3 证书",
                    "file": "cmmi3.png",
                    "expiry": "2027-08-01",
                },
                {
                    "id": "ISO9001",
                    "name": "ISO9001 认证",
                    "file": "iso9001.png",
                    "expiry": "2029-03-01",
                },
            ],
        },
    )


def test_finalize_produces_valid_docx(tmp_path):
    ws = BidWorkspace("asm1", root=tmp_path)
    _seed(ws)
    out = assemble.finalize(ws)
    assert out.exists() and out.name == "投标文件.docx"
    # resolve replaced placeholders into workspace/final
    final = ws.path("workspace/final/s1.md").read_text(encoding="utf-8")
    assert "{{" not in final and "测试科技有限公司" in final and "（见附件1）" in final
    assert ws.path("workspace/attachments.json").exists()
    # docx is a real, openable Word file with the cover + section headings
    doc = Document(str(out))
    heads = [
        p.text for p in doc.paragraphs if p.style.name.startswith(("Heading", "Title"))
    ]
    assert any("总述" in h for h in heads) and any("附件" in h for h in heads)


def test_finalize_dangling_placeholder_raises(tmp_path):
    ws = BidWorkspace("asm2", root=tmp_path)
    _seed(ws)
    # reference a qual id absent from the knowledge base -> resolve hard gate
    ws.path("workspace/sections/s1.md").write_text(
        "{{bidder}} 持有 {{qual:GHOST}}。", encoding="utf-8"
    )
    with pytest.raises(BidPipelineError):
        assemble.finalize(ws)


def test_finalize_injects_section_titles_from_outline(tmp_path):
    # Regression for the title-dedup change: bodies carry no heading now, so the
    # assembler MUST add each section's heading from outline[*]["title"]. Without
    # that, title-less bodies would silently lose their headings on export.
    ws = BidWorkspace("asm3", root=tmp_path)
    _seed(ws)
    out = assemble.finalize(ws)
    doc = Document(str(out))
    heads = [
        p.text for p in doc.paragraphs if p.style.name.startswith(("Heading", "Title"))
    ]
    # The two outline titles must each appear as a heading.
    assert any("总述" in h for h in heads)
    assert any("质量" in h for h in heads)
    # And the bodies must still be present (post-resolve).
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "CMMI3" in body and "ISO9001" in body
