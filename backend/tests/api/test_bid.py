# SPDX-License-Identifier: Apache-2.0
from unittest.mock import AsyncMock, patch

import pytest

from app.core.config import settings
from app.core.security import create_access_token, get_password_hash
from app.models.user import User
from app.services.bid.workspace import BidWorkspace


def _second_user(test_db) -> tuple[User, str]:
    # Field set mirrors tests/conftest.py::test_user (User requires user_name +
    # password_hash NOT NULL; is_active/git_info match the existing fixture).
    u = User(
        user_name="user2",
        password_hash=get_password_hash("user2password123"),
        email="u2@x.com",
        is_active=True,
        git_info=None,
    )
    test_db.add(u)
    test_db.commit()
    test_db.refresh(u)
    return u, create_access_token(data={"sub": u.user_name})


def test_bid_crud_and_parse(test_client, test_token, tmp_path, monkeypatch):
    # Redirect per-project workspaces to a writable temp root (default points at
    # /data/bid-workspaces which is not writable in the test environment).
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post(
        "/api/bid/projects", json={"title": "智慧园区"}, headers=h
    ).json()["id"]
    assert any(
        p["id"] == pid for p in test_client.get("/api/bid/projects", headers=h).json()
    )

    # 拆标 is backgrounded: the endpoint returns 'parsing' and schedules
    # launch_parse (the client polls project status). Patch launch_parse so the
    # background task does not actually run in the test.
    with (
        patch("app.api.endpoints.bid.launch_parse") as launch,
        patch(
            "app.api.endpoints.bid.resolve_project_model",
            return_value=("m", {"api_key": "k"}),
        ),
    ):
        r = test_client.post(
            f"/api/bid/projects/{pid}/parse",
            json={"tender_text": "招标正文"},
            headers=h,
        )
    assert r.status_code == 200 and r.json()["status"] == "parsing"
    launch.assert_called_once()
    # tender text is written to the blackboard for the background task
    ref = test_client.get(f"/api/bid/projects/{pid}", headers=h).json()["workspace_ref"]
    ws = BidWorkspace(ref, root=tmp_path)
    assert ws.path("inputs/final/tender.txt").read_text(encoding="utf-8") == "招标正文"
    # project is locked in 'parsing' until the background task completes
    assert (
        test_client.get(f"/api/bid/projects/{pid}", headers=h).json()["status"]
        == "parsing"
    )


def test_bid_ownership_isolation(
    test_client, test_token, test_db, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post("/api/bid/projects", json={"title": "A"}, headers=h).json()[
        "id"
    ]
    _u2, tok2 = _second_user(test_db)
    r = test_client.get(
        f"/api/bid/projects/{pid}", headers={"Authorization": f"Bearer {tok2}"}
    )
    assert r.status_code == 404


# NOTE: parse-failure lock release now lives in parse_pipeline._run_parse
# (the background task); see test_parse_pipeline.test_run_parse_marks_parse_failed_on_error.


def test_outline_build_read_edit_coverage(
    test_client, test_token, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post("/api/bid/projects", json={"title": "O"}, headers=h).json()[
        "id"
    ]

    # Pre-seed tender.json into the project workspace. workspace_ref is now
    # exposed via the project response so the test can locate the blackboard.
    ws_ref = test_client.get(f"/api/bid/projects/{pid}", headers=h).json()[
        "workspace_ref"
    ]
    ws = BidWorkspace(ws_ref)
    ws.write_json(
        "workspace/tender.json",
        {
            "scoring": [{"id": "S1"}],
            "mandatory_clauses": [{"id": "V1", "veto": True}],
        },
    )

    async def _build(_ws):
        _ws.write_json(
            "workspace/outline.json",
            {"sections": [{"id": "a", "covers": ["S1"]}], "volumes": []},
        )
        return _ws.read_json("workspace/outline.json")

    with patch(
        "app.api.endpoints.bid.build_outline_for_project",
        new=AsyncMock(side_effect=_build),
    ):
        r = test_client.post(f"/api/bid/projects/{pid}/outline", headers=h)
    assert r.status_code == 200
    # Uncovered items carry {id, text}; V1 has no text so text is "".
    uc = r.json()["coverage"]["uncovered_clauses"]
    assert [x["id"] for x in uc] == ["V1"]  # V1 not covered

    # Human edit writeback: add V1 to coverage.
    edited = {"sections": [{"id": "a", "covers": ["S1", "V1"]}], "volumes": []}
    r2 = test_client.put(
        f"/api/bid/projects/{pid}/outline", json={"outline": edited}, headers=h
    )
    assert r2.json()["coverage"]["uncovered_clauses"] == []


def test_package_declaration(test_client, test_token, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post("/api/bid/projects", json={"title": "P"}, headers=h).json()[
        "id"
    ]
    r = test_client.post(
        f"/api/bid/projects/{pid}/package", json={"package": "包件二"}, headers=h
    )
    assert r.status_code == 200 and r.json()["status"] == "declared"
    ws_ref = test_client.get(f"/api/bid/projects/{pid}", headers=h).json()[
        "workspace_ref"
    ]
    ws = BidWorkspace(ws_ref)
    assert ws.read_json("corpus/bid_config.json")["package"] == "包件二"


def test_materials_kb_quals_attachments_and_complete(
    test_client, test_token, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post("/api/bid/projects", json={"title": "M"}, headers=h).json()[
        "id"
    ]

    # knowledge base put/get
    kb = {"knowledge_base": {"bidder_knowledge_base": {"cases": [{"name": "案例甲"}]}}}
    assert (
        test_client.put(
            f"/api/bid/projects/{pid}/materials/knowledge-base", json=kb, headers=h
        ).status_code
        == 200
    )
    got = test_client.get(
        f"/api/bid/projects/{pid}/materials/knowledge-base", headers=h
    ).json()
    assert (
        got["knowledge_base"]["bidder_knowledge_base"]["cases"][0]["name"] == "案例甲"
    )

    # bad shape -> 400
    bad = {"knowledge_base": {"nope": 1}}
    assert (
        test_client.put(
            f"/api/bid/projects/{pid}/materials/knowledge-base", json=bad, headers=h
        ).status_code
        == 400
    )

    # qualifications
    q = {
        "qualifications": {
            "company": "好大一家",
            "items": [{"id": "Q1", "name": "ISO9001"}],
        }
    }
    assert (
        test_client.put(
            f"/api/bid/projects/{pid}/materials/qualifications", json=q, headers=h
        ).status_code
        == 200
    )

    # attachment upload + list
    r = test_client.post(
        f"/api/bid/projects/{pid}/materials/attachments",
        files={"file": ("iso9001.pdf", b"PDFDATA", "application/pdf")},
        headers=h,
    )
    assert r.status_code == 200 and r.json() == {
        "name": "iso9001.pdf",
        "size": 7,
        "stats": None,
    }
    lst = test_client.get(
        f"/api/bid/projects/{pid}/materials/attachments", headers=h
    ).json()
    assert lst["items"] == [{"name": "iso9001.pdf", "size": 7, "stats": None}]

    # complete -> phase advances
    assert (
        test_client.post(
            f"/api/bid/projects/{pid}/materials/complete", headers=h
        ).json()["status"]
        == "materials_done"
    )
    assert (
        test_client.get(f"/api/bid/projects/{pid}", headers=h).json()["current_phase"]
        == 4
    )


def test_materials_kb_missing_returns_409(
    test_client, test_token, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post("/api/bid/projects", json={"title": "N"}, headers=h).json()[
        "id"
    ]
    assert (
        test_client.get(
            f"/api/bid/projects/{pid}/materials/knowledge-base", headers=h
        ).status_code
        == 409
    )


def test_draft_trigger_and_status(test_client, test_token, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post("/api/bid/projects", json={"title": "D"}, headers=h).json()[
        "id"
    ]
    ref = test_client.get(f"/api/bid/projects/{pid}", headers=h).json()["workspace_ref"]

    ws = BidWorkspace(ref, root=tmp_path)
    # no outline -> 409
    assert (
        test_client.post(f"/api/bid/projects/{pid}/draft", headers=h).status_code == 409
    )
    ws.write_json("workspace/outline.json", {"sections": [{"id": "s1", "covers": []}]})

    with (
        patch("app.api.endpoints.bid.launch_drafting") as launch,
        patch(
            "app.api.endpoints.bid.resolve_project_model",
            return_value=("m", {"api_key": "k"}),
        ),
    ):
        r = test_client.post(f"/api/bid/projects/{pid}/draft", headers=h)
    assert r.status_code == 200 and r.json()["status"] == "drafting"
    launch.assert_called_once()
    # placeholder status readable
    assert (
        test_client.get(f"/api/bid/projects/{pid}/draft/status", headers=h).json()[
            "finished"
        ]
        is False
    )

    # pre-seed a section artifact -> sections list + read single section
    ws.path("workspace/sections").mkdir(parents=True, exist_ok=True)
    ws.path("workspace/sections/s1.md").write_text("正文一", encoding="utf-8")
    assert any(
        i["id"] == "s1"
        for i in test_client.get(f"/api/bid/projects/{pid}/sections", headers=h).json()[
            "items"
        ]
    )
    assert (
        test_client.get(f"/api/bid/projects/{pid}/sections/s1", headers=h).json()[
            "content"
        ]
        == "正文一"
    )


def test_review_redraft_accept_complete(test_client, test_token, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post("/api/bid/projects", json={"title": "R"}, headers=h).json()[
        "id"
    ]
    ref = test_client.get(f"/api/bid/projects/{pid}", headers=h).json()["workspace_ref"]
    ws = BidWorkspace(ref, root=tmp_path)
    ws.write_json("workspace/outline.json", {"sections": [{"id": "s1", "covers": []}]})

    # section not in outline -> 404
    assert (
        test_client.post(
            f"/api/bid/projects/{pid}/sections/zzz/redraft", json={}, headers=h
        ).status_code
        == 404
    )

    with (
        patch("app.api.endpoints.bid.launch_redraft") as launch,
        patch(
            "app.api.endpoints.bid.resolve_project_model",
            return_value=("m", {"api_key": "k"}),
        ),
    ):
        r = test_client.post(
            f"/api/bid/projects/{pid}/sections/s1/redraft",
            json={"instruction": "更简洁"},
            headers=h,
        )
    assert r.status_code == 200 and r.json()["status"] == "drafting"
    launch.assert_called_once()

    assert (
        test_client.post(
            f"/api/bid/projects/{pid}/sections/s1/accept", headers=h
        ).json()["status"]
        == "accepted"
    )
    assert test_client.get(f"/api/bid/projects/{pid}/review/status", headers=h).json()[
        "accepted"
    ] == {"s1": True}
    assert (
        test_client.post(f"/api/bid/projects/{pid}/review/complete", headers=h).json()[
            "status"
        ]
        == "review_done"
    )
    assert (
        test_client.get(f"/api/bid/projects/{pid}", headers=h).json()["current_phase"]
        == 6
    )


def test_save_section_cas_conflict_and_accept_clear(
    test_client, test_token, tmp_path, monkeypatch
):
    # PUT /sections/{sid}/content: stale base_version -> 409; correct version ->
    # 200, returns new version, and clears any prior "accepted" flag for the section.
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post(
        "/api/bid/projects", json={"title": "save"}, headers=h
    ).json()["id"]
    ref = test_client.get(f"/api/bid/projects/{pid}", headers=h).json()["workspace_ref"]
    ws = BidWorkspace(ref, root=tmp_path)
    ws.path("workspace/sections").mkdir(parents=True, exist_ok=True)
    ws.path("workspace/sections/s1.md").write_text("初稿", encoding="utf-8")
    # mark the section accepted up front so we can assert editing clears it
    ws.write_json("workspace/_review_status.json", {"accepted": {"s1": True}})

    # GET returns the current version (sha256 of file bytes)
    got = test_client.get(f"/api/bid/projects/{pid}/sections/s1", headers=h).json()
    assert got["content"] == "初稿"
    assert got["version"]

    # save with a stale base_version -> 409
    r = test_client.put(
        f"/api/bid/projects/{pid}/sections/s1/content",
        json={"content": "edited", "base_version": "deadbeef"},
        headers=h,
    )
    assert r.status_code == 409

    # save with the correct base_version -> 200 + new version
    r2 = test_client.put(
        f"/api/bid/projects/{pid}/sections/s1/content",
        json={"content": "edited", "base_version": got["version"]},
        headers=h,
    )
    assert r2.status_code == 200
    assert r2.json()["version"] != got["version"]
    # content persisted
    assert (
        test_client.get(f"/api/bid/projects/{pid}/sections/s1", headers=h).json()[
            "content"
        ]
        == "edited"
    )
    # editing cleared the accepted flag
    assert (
        "s1"
        not in test_client.get(
            f"/api/bid/projects/{pid}/review/status", headers=h
        ).json()["accepted"]
    )


def test_redraft_range_stale_version_returns_409(
    test_client, test_token, tmp_path, monkeypatch
):
    # POST /sections/{sid}/redraft-range: CAS on base_version -> 409 when stale.
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post("/api/bid/projects", json={"title": "rr"}, headers=h).json()[
        "id"
    ]
    ref = test_client.get(f"/api/bid/projects/{pid}", headers=h).json()["workspace_ref"]
    ws = BidWorkspace(ref, root=tmp_path)
    ws.path("workspace/sections").mkdir(parents=True, exist_ok=True)
    ws.path("workspace/sections/s1.md").write_text("A\nB\nC", encoding="utf-8")

    r = test_client.post(
        f"/api/bid/projects/{pid}/sections/s1/redraft-range",
        json={
            "start_line": 1,
            "end_line": 1,
            "instruction": None,
            "base_version": "stale",
        },
        headers=h,
    )
    assert r.status_code == 409


def test_audit_finalize_download(test_client, test_token, tmp_path, monkeypatch):
    from app.core.config import settings

    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post("/api/bid/projects", json={"title": "F"}, headers=h).json()[
        "id"
    ]
    ref = test_client.get(f"/api/bid/projects/{pid}", headers=h).json()["workspace_ref"]

    from app.services.bid.workspace import BidWorkspace

    ws = BidWorkspace(ref, root=tmp_path)
    ws.write_json(
        "workspace/tender.json",
        {
            "project": {
                "name": "P",
                "id": "1",
                "budget": "100",
                "bid_deadline": "2026-12-31T00:00:00",
            },
            "submission_rules": {"blind_bid": False},
        },
    )
    ws.write_json(
        "workspace/outline.json",
        {"sections": [{"id": "s1", "title": "总述", "covers": []}]},
    )
    ws.path("workspace/sections").mkdir(parents=True, exist_ok=True)
    ws.path("workspace/sections/s1.md").write_text(
        "# 总述\n\n由{{bidder}}实施{{qual:CMMI3}}。", encoding="utf-8"
    )
    ws.write_json(
        "corpus/qualifications.json",
        {
            "company": "测试公司",
            "items": [
                {
                    "id": "CMMI3",
                    "name": "CMMI3 证书",
                    "file": "cmmi3.png",
                    "expiry": "2027-08-01",
                }
            ],
        },
    )

    # download before finalize -> 404
    assert (
        test_client.get(f"/api/bid/projects/{pid}/download", headers=h).status_code
        == 404
    )
    # audit report before audit -> 404
    assert (
        test_client.get(f"/api/bid/projects/{pid}/audit/report", headers=h).status_code
        == 404
    )

    # audit does NOT advance phase (rework gate)
    r = test_client.post(f"/api/bid/projects/{pid}/audit", headers=h)
    assert r.status_code == 200 and r.json()["verdict"] in (
        "PASS",
        "NEED_FIX",
        "NEED_FIX_VETO",
    )
    assert (
        test_client.get(f"/api/bid/projects/{pid}", headers=h).json()["current_phase"]
        == 1
    )
    assert (
        test_client.get(f"/api/bid/projects/{pid}/audit/report", headers=h).json()[
            "verdict"
        ]
        == r.json()["verdict"]
    )

    # finalize -> docx + phase advanced to 7 (set_phase_done(6))
    assert (
        test_client.post(f"/api/bid/projects/{pid}/finalize", headers=h).json()[
            "status"
        ]
        == "finalized"
    )
    assert (
        test_client.get(f"/api/bid/projects/{pid}", headers=h).json()["current_phase"]
        == 7
    )

    # download -> 200 docx
    dl = test_client.get(f"/api/bid/projects/{pid}/download", headers=h)
    assert dl.status_code == 200
    assert dl.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert dl.content[:2] == b"PK"  # docx is a zip


def test_audit_verify_folds_fidelity_verdicts(
    test_client, test_token, tmp_path, monkeypatch
):
    from unittest.mock import AsyncMock, patch

    from app.core.config import settings

    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post("/api/bid/projects", json={"title": "V"}, headers=h).json()[
        "id"
    ]
    ref = test_client.get(f"/api/bid/projects/{pid}", headers=h).json()["workspace_ref"]

    from app.services.bid.workspace import BidWorkspace

    ws = BidWorkspace(ref, root=tmp_path)
    ws.write_json(
        "workspace/tender.json",
        {
            "project": {
                "name": "P",
                "id": "1",
                "budget": "100",
                "bid_deadline": "2026-12-31T00:00:00",
            },
            "submission_rules": {"blind_bid": False},
        },
    )
    ws.write_json(
        "workspace/outline.json",
        {"sections": [{"id": "s1", "title": "总述", "covers": []}]},
    )
    ws.path("workspace/sections").mkdir(parents=True, exist_ok=True)
    ws.path("workspace/sections/s1.md").write_text("# 总述\n正文。", encoding="utf-8")
    ws.write_json("corpus/qualifications.json", {"company": "测试公司", "items": []})

    # no tasks yet, no report -> 409
    assert (
        test_client.post(f"/api/bid/projects/{pid}/audit/verify", headers=h).status_code
        == 409
    )

    # seed a fidelity task (as Plan 11 /audit first pass would)
    ws.write_json(
        "workspace/_fidelity_tasks.json",
        {
            "stage": "audit",
            "tasks": [
                {
                    "id": "SF-0",
                    "type": "clause_faithfulness",
                    "claim": "响应全部技术要求",
                    "candidate_sources": [],
                }
            ],
        },
    )
    with (
        patch(
            "app.api.endpoints.bid.call_fact_checker",
            new=AsyncMock(
                return_value=[
                    {
                        "id": "SF-0",
                        "verdict": "uncertain",
                        "reason": "证据不足",
                        "severity": "medium",
                    }
                ]
            ),
        ) as fc,
        patch(
            "app.api.endpoints.bid.resolve_project_model",
            return_value=("m", {"api_key": "k"}),
        ),
    ):
        r = test_client.post(f"/api/bid/projects/{pid}/audit/verify", headers=h)
    assert r.status_code == 200 and r.json()["verdict"] in (
        "PASS",
        "NEED_FIX",
        "NEED_FIX_VETO",
    )
    fc.assert_awaited_once()
    assert ws.path("workspace/_fidelity_verdicts.json").exists()
    # verify does NOT advance phase
    assert (
        test_client.get(f"/api/bid/projects/{pid}", headers=h).json()["current_phase"]
        == 1
    )


def test_briefs_roundtrip_api(test_client, test_token, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post(
        "/api/bid/projects", json={"title": "briefs项目"}, headers=h
    ).json()["id"]

    # empty default, no 4xx noise on first load
    r = test_client.get(f"/api/bid/projects/{pid}/materials/briefs", headers=h)
    assert r.status_code == 200 and r.json() == {"briefs": {}, "materials": []}

    # Backing file must exist on disk or the read-time reconcile drops the entry.
    test_client.post(
        f"/api/bid/projects/{pid}/materials/attachments",
        files={"file": ("a.pdf", b"PDF", "application/pdf")},
        headers=h,
    )
    doc = {
        "briefs": {"s1": {"style": "专业", "requirements": "写清楚", "tags": ["核心"]}},
        "materials": [
            {"id": "m1", "name": "a.pdf", "size": 3, "linkedNodeIds": ["s1"]}
        ],
    }
    r = test_client.put(
        f"/api/bid/projects/{pid}/materials/briefs", json=doc, headers=h
    )
    # PUT writes the doc verbatim (no reconcile on write — only read aligns).
    assert r.status_code == 200 and r.json() == doc
    r = test_client.get(f"/api/bid/projects/{pid}/materials/briefs", headers=h)
    # GET reconciles against on-disk attachments: a.pdf exists, so it survives.
    got = r.json()
    assert got["briefs"] == doc["briefs"]
    assert [m["name"] for m in got["materials"]] == ["a.pdf"]

    # invalid shape -> 400
    r = test_client.put(
        f"/api/bid/projects/{pid}/materials/briefs",
        json={"briefs": {"s1": "oops"}},
        headers=h,
    )
    assert r.status_code == 400


def test_extract_text_api(test_client, test_token):
    h = {"Authorization": f"Bearer {test_token}"}
    body = ("招标文件正文，第一章 项目概述。" * 10).encode("utf-8")
    r = test_client.post(
        "/api/bid/extract-text",
        files={"file": ("标书.txt", body, "text/plain")},
        headers=h,
    )
    assert r.status_code == 200
    data = r.json()
    assert data["name"] == "标书.txt" and "项目概述" in data["text"]

    r = test_client.post(
        "/api/bid/extract-text",
        files={"file": ("x.doc", b"binary", "application/msword")},
        headers=h,
    )
    assert r.status_code == 422


def test_llm_log_endpoint_empty(test_client, test_token, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post(
        "/api/bid/projects", json={"title": "log"}, headers=h
    ).json()["id"]
    # empty first (no calls recorded yet)
    assert (
        test_client.get(f"/api/bid/projects/{pid}/llm-log", headers=h).json()["items"]
        == []
    )


def test_parse_stage_endpoint(test_client, test_token, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post("/api/bid/projects", json={"title": "ps"}, headers=h).json()[
        "id"
    ]
    # idle before any parse runs
    assert (
        test_client.get(f"/api/bid/projects/{pid}/parse-stage", headers=h).json()[
            "stage"
        ]
        == "idle"
    )


def test_upload_attachment_returns_stats(
    test_client, test_token, tmp_path, monkeypatch
):
    import io

    from docx import Document

    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post("/api/bid/projects", json={"title": "M"}, headers=h).json()[
        "id"
    ]

    doc = Document()
    doc.add_paragraph("投标人具备完善能力。" * 8)
    doc.add_table(rows=2, cols=2)
    buf = io.BytesIO()
    doc.save(buf)

    r = test_client.post(
        f"/api/bid/projects/{pid}/materials/attachments",
        files={
            "file": (
                "cap.docx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=h,
    )
    assert r.status_code == 200
    body = r.json()
    assert body["name"] == "cap.docx"
    assert body["stats"]["tables"] == 1
    assert body["stats"]["chars"] >= 30


def test_get_grounding_maps_nodes_to_covered_items(tmp_path, monkeypatch):
    # Service-level: drive the endpoint function directly with a stubbed project.
    from app.api.endpoints import bid as bid_ep
    from app.services.bid.workspace import BidWorkspace

    ws = BidWorkspace("grounding-1", root=tmp_path)
    ws.write_json(
        "workspace/tender.json",
        {
            "scoring": [{"id": "T1", "item": "方案", "weight": 20}],
            "mandatory_clauses": [{"id": "M1", "veto": True, "clause": "有效期"}],
        },
    )
    ws.write_json(
        "workspace/outline.json",
        {
            "sections": [
                {
                    "id": "s1",
                    "title": "总述",
                    "covers": ["T1"],
                    "children": [{"id": "s1.1", "covers": ["M1"]}],
                },
                {"id": "s2", "title": "空节", "covers": []},
            ]
        },
    )

    items = bid_ep._grounding_items(ws)  # helper extracted in Step 3

    assert set(items.keys()) == {"s1", "s1.1"}  # s2 has no grounding -> omitted
    assert [s["id"] for s in items["s1"]["scoring"]] == ["T1"]
    assert [c["id"] for c in items["s1.1"]["clauses"]] == ["M1"]


@pytest.mark.asyncio
async def test_generate_briefs_endpoint_returns_briefs(tmp_path, monkeypatch):
    from unittest.mock import AsyncMock

    from app.api.endpoints import bid as bid_ep

    async def fake_gen(ws, **kw):
        return {
            "s1": {
                "requirements": "写详细",
                "emphasis": "亮点",
                "wordMin": "800",
                "wordMax": "1500",
                "needFigure": "否",
                "importance": "中",
            }
        }

    monkeypatch.setattr(
        bid_ep.brief_pipeline, "generate_briefs", AsyncMock(side_effect=fake_gen)
    )
    out = await bid_ep._generate_briefs_impl(
        ws=bid_ep.BidWorkspace("gen-ep", root=tmp_path),
        model="m",
        model_config=None,
        node_ids=["s1"],
        project_id=1,
        user_id=2,
    )
    assert out["s1"]["requirements"] == "写详细"


def test_brief_status_endpoint_reads_status(tmp_path):
    from app.api.endpoints import bid as bid_ep
    from app.services.bid import brief_pipeline as bp
    from app.services.bid.workspace import BidWorkspace

    ws = BidWorkspace("brief-ep", root=tmp_path)
    bp.init_brief_status(ws, ["s1"])
    st = bid_ep.brief_pipeline.read_brief_status(ws)
    assert st["nodes"] == {"s1": "pending"}
    # endpoints must be registered
    assert hasattr(bid_ep, "auto_generate_briefs")
    assert hasattr(bid_ep, "brief_status")


def test_ingest_endpoint_reingests_all(tmp_path, monkeypatch):
    from app.api.endpoints import bid as bid_ep
    from app.services.bid.workspace import BidWorkspace

    assert hasattr(bid_ep, "ingest_materials")

    ws = BidWorkspace("ingest-ep", root=tmp_path)
    ws.path("corpus/attachments/a.txt").parent.mkdir(parents=True, exist_ok=True)
    ws.path("corpus/attachments/a.txt").write_bytes(b"hello world")
    monkeypatch.setattr(
        bid_ep.materials_store, "extract_text", lambda n, c: "hello world"
    )
    out = bid_ep.materials_store.ingest_all(ws)
    assert out == {"a.txt": "ok"}


def test_outline_stage3_read_edit_and_differs(
    test_client, test_token, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post("/api/bid/projects", json={"title": "S3"}, headers=h).json()[
        "id"
    ]
    ws_ref = test_client.get(f"/api/bid/projects/{pid}", headers=h).json()[
        "workspace_ref"
    ]
    ws = BidWorkspace(ws_ref)
    ws.write_json(
        "workspace/outline.json",
        {"sections": [{"id": "a", "title": "第一章", "covers": []}], "volumes": []},
    )
    # The stage-1 GET /outline at the end computes coverage, which needs tender.
    ws.write_json(
        "workspace/tender.json",
        {"scoring": [], "mandatory_clauses": []},
    )

    # GET lazily copies stage1 -> stage3; fresh copy does not differ.
    r = test_client.get(f"/api/bid/projects/{pid}/outline/stage3", headers=h)
    assert r.status_code == 200
    assert r.json()["outline"]["sections"][0]["id"] == "a"
    assert r.json()["differs_from_stage1"] is False

    # PUT an edited stage3 -> differs; stage1 untouched.
    edited = {
        "sections": [{"id": "a", "title": "改过的第一章", "covers": []}],
        "volumes": [],
    }
    r2 = test_client.put(
        f"/api/bid/projects/{pid}/outline/stage3", json={"outline": edited}, headers=h
    )
    assert r2.status_code == 200 and r2.json()["differs_from_stage1"] is True
    assert (
        test_client.get(f"/api/bid/projects/{pid}/outline", headers=h).json()[
            "outline"
        ]["sections"][0]["title"]
        == "第一章"
    )


def test_draft_resets_stage3_to_stage1(test_client, test_token, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "BID_WORKSPACE_ROOT", str(tmp_path))
    h = {"Authorization": f"Bearer {test_token}"}
    pid = test_client.post("/api/bid/projects", json={"title": "R"}, headers=h).json()[
        "id"
    ]
    ref = test_client.get(f"/api/bid/projects/{pid}", headers=h).json()["workspace_ref"]
    ws = BidWorkspace(ref, root=tmp_path)
    ws.write_json(
        "workspace/outline.json",
        {"sections": [{"id": "s1", "title": "第一章", "covers": []}], "volumes": []},
    )
    # A diverged stage3 (as if the user edited it in stage 3).
    ws.write_json(
        "workspace/outline_stage3.json",
        {"sections": [{"id": "s1", "title": "被改过", "covers": []}], "volumes": []},
    )

    with (
        patch("app.api.endpoints.bid.launch_drafting"),
        patch(
            "app.api.endpoints.bid.resolve_project_model",
            return_value=("m", {"api_key": "k"}),
        ),
    ):
        r = test_client.post(f"/api/bid/projects/{pid}/draft", headers=h)
    assert r.status_code == 200
    # Full re-draft overwrote stage3 back to stage1.
    assert (
        ws.read_json("workspace/outline_stage3.json")["sections"][0]["title"]
        == "第一章"
    )
